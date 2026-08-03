from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt

from models import CoverLetterDraft, Profile, ResumeDraft

_CONTENT_WIDTH = Inches(7.0)  # letter width minus 0.75" margins


def _base_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    return doc


def _entry_header(doc: Document, title: str, organization: str, date_range: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.tab_stops.add_tab_stop(_CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    run = paragraph.add_run(title)
    run.bold = True
    if organization:
        paragraph.add_run(f" — {organization}")
    if date_range:
        date_run = paragraph.add_run(f"\t{date_range}")
        date_run.italic = True


def resume_docx(draft: ResumeDraft, profile: Profile | None = None) -> bytes:
    profile = profile or Profile()
    doc = _base_document()
    doc.add_heading(profile.name or "Your Name", 0)
    contact = profile.contact_line()
    if contact:
        contact_paragraph = doc.add_paragraph(contact)
        contact_paragraph.paragraph_format.space_after = Pt(8)
    doc.add_paragraph(draft.summary)
    doc.add_heading("Skills", level=1)
    doc.add_paragraph(" • ".join(draft.skills))
    for section in draft.sections:
        doc.add_heading(section.heading, level=1)
        for entry in section.entries:
            _entry_header(doc, entry.title, entry.organization, entry.date_range)
            for bullet in entry.bullets:
                doc.add_paragraph(bullet.text, style="List Bullet")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def cover_letter_docx(draft: CoverLetterDraft) -> bytes:
    doc = _base_document()
    doc.add_paragraph(draft.salutation)
    for paragraph in draft.paragraphs:
        doc.add_paragraph(paragraph)
    doc.add_paragraph(draft.closing)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
